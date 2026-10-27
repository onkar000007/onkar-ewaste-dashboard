import os
import sys
import zipfile
import re
import html
import docx
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
import openpyxl
import xml.etree.ElementTree as ET

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.abspath('assets/assignment-files')
PREVIEWS_DIR = os.path.abspath('assets/previews')
MEDIA_DIR = os.path.join(PREVIEWS_DIR, 'media')

os.makedirs(PREVIEWS_DIR, exist_ok=True)
os.makedirs(MEDIA_DIR, exist_ok=True)

COMMON_CSS = """
:root {
  --bg: #f8fafc;
  --surface: #ffffff;
  --primary: #2563eb;
  --text-main: #0f172a;
  --text-muted: #64748b;
  --border: #e2e8f0;
  --tag-bg: #eff6ff;
  --tag-text: #1d4ed8;
}
* { box-sizing: border-box; margin: 0; padding: 0; }
body {
  font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
  background: var(--bg);
  color: var(--text-main);
  line-height: 1.65;
  padding: 18px;
}
.preview-container {
  max-width: 960px;
  margin: 0 auto;
  background: var(--surface);
  border-radius: 14px;
  box-shadow: 0 4px 20px -2px rgba(15, 23, 42, 0.08);
  border: 1px solid var(--border);
  overflow: hidden;
}
.doc-header {
  padding: 22px 30px 16px;
  border-bottom: 1px solid var(--border);
  background: linear-gradient(to bottom, #ffffff, #fbfcfe);
}
.badge-pill {
  display: inline-block;
  padding: 4px 10px;
  background: var(--tag-bg);
  color: var(--tag-text);
  border-radius: 9999px;
  font-size: 11px;
  font-weight: 700;
  letter-spacing: 0.05em;
  text-transform: uppercase;
  margin-bottom: 8px;
}
.doc-title {
  font-size: 22px;
  font-weight: 800;
  color: var(--text-main);
  line-height: 1.3;
  margin-bottom: 8px;
}
.doc-meta {
  display: flex;
  flex-wrap: wrap;
  gap: 14px;
  font-size: 13px;
  color: var(--text-muted);
}
.meta-item strong { color: var(--text-main); }
.actions-bar {
  display: flex;
  align-items: center;
  justify-content: space-between;
  padding: 10px 30px;
  background: #f8fafc;
  border-bottom: 1px solid var(--border);
}
.btn-action {
  display: inline-flex;
  align-items: center;
  gap: 6px;
  background: #ffffff;
  border: 1px solid var(--border);
  padding: 6px 12px;
  border-radius: 8px;
  font-size: 12px;
  font-weight: 600;
  color: var(--text-main);
  text-decoration: none;
  cursor: pointer;
}
.doc-body {
  padding: 28px 30px;
  font-size: 14.5px;
}
.doc-body h1 {
  font-size: 19px;
  font-weight: 700;
  margin: 26px 0 10px;
  color: #1e293b;
  border-bottom: 2px solid #e2e8f0;
  padding-bottom: 6px;
}
.doc-body h2 {
  font-size: 16px;
  font-weight: 700;
  margin: 20px 0 8px;
  color: #334155;
}
.doc-body h3 {
  font-size: 14.5px;
  font-weight: 600;
  margin: 14px 0 6px;
  color: #475569;
}
.doc-body p {
  margin-bottom: 12px;
  color: #334155;
  text-align: justify;
}
.doc-body ul, .doc-body ol {
  margin: 10px 0 14px 22px;
  color: #334155;
}
.doc-body li { margin-bottom: 5px; }
.doc-body img {
  max-width: 100%;
  height: auto;
  border-radius: 8px;
  display: block;
  margin: 16px auto;
  border: 1px solid var(--border);
}
.fig-caption {
  text-align: center;
  font-size: 12px;
  color: var(--text-muted);
  font-style: italic;
  margin-top: -8px;
  margin-bottom: 14px;
}
.table-wrapper {
  overflow-x: auto;
  margin: 16px 0 20px;
  border-radius: 8px;
  border: 1px solid var(--border);
}
table {
  width: 100%;
  border-collapse: collapse;
  font-size: 13px;
  text-align: left;
}
th {
  background: #f1f5f9;
  color: #1e293b;
  font-weight: 600;
  padding: 9px 12px;
  border-bottom: 2px solid var(--border);
}
td {
  padding: 8px 12px;
  border-bottom: 1px solid var(--border);
  color: #334155;
}
tr:last-child td { border-bottom: none; }
tr:nth-child(even) td { background: #fafbfc; }

/* Tabs for XLSX */
.sheet-tabs {
  display: flex;
  gap: 6px;
  overflow-x: auto;
  padding: 12px 30px;
  background: #f1f5f9;
  border-bottom: 1px solid var(--border);
}
.sheet-tab {
  padding: 6px 14px;
  background: #e2e8f0;
  border-radius: 8px;
  font-size: 12px;
  font-weight: 600;
  cursor: pointer;
  white-space: nowrap;
  border: none;
  color: #334155;
}
.sheet-tab.active {
  background: var(--primary);
  color: white;
}
.sheet-content { display: none; padding: 24px 30px; }
.sheet-content.active { display: block; }
"""

def extract_media_from_zip(zip_path, prefix='word/media/'):
    saved = {}
    with zipfile.ZipFile(zip_path) as z:
        for name in z.namelist():
            if name.startswith(prefix):
                base_name = os.path.basename(name)
                doc_prefix = os.path.splitext(os.path.basename(zip_path))[0][:15].strip()
                doc_prefix = re.sub(r'[^a-zA-Z0-9_]', '_', doc_prefix)
                unique_name = f"{doc_prefix}_{base_name}"
                out_path = os.path.join(MEDIA_DIR, unique_name)
                with open(out_path, 'wb') as f_out:
                    f_out.write(z.read(name))
                saved[base_name] = unique_name
    return saved

def convert_docx(filename):
    full_path = os.path.join(BASE_DIR, filename)
    if not os.path.exists(full_path):
        return
    media_map = extract_media_from_zip(full_path, 'word/media/')
    doc = docx.Document(full_path)
    body_html = []
    
    def format_para(p):
        p_html = []
        for run in p.runs:
            t = html.escape(run.text)
            if not t: continue
            if run.bold and run.italic: t = f"<strong><em>{t}</em></strong>"
            elif run.bold: t = f"<strong>{t}</strong>"
            elif run.italic: t = f"<em>{t}</em>"
            p_html.append(t)
        return "".join(p_html)

    def get_para_images(p):
        imgs = []
        xml = p._element.xml
        rids = re.findall(r'embed=[\'"]([^\'"]+)[\'"]', xml)
        for rid in rids:
            part = doc.part.related_parts.get(rid)
            if part and hasattr(part, 'partname'):
                pname = os.path.basename(part.partname)
                if pname in media_map:
                    imgs.append(media_map[pname])
        return imgs

    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            p = Paragraph(child, doc)
            text = p.text.strip()
            for img in get_para_images(p):
                body_html.append(f'<img src="media/{img}" alt="Document Figure">')
            if not text: continue
            style_name = p.style.name if p.style else ''
            if 'Heading 1' in style_name or (len(text) < 70 and p.runs and p.runs[0].bold and text.isupper() and len(text) > 4):
                body_html.append(f'<h1>{html.escape(text)}</h1>')
            elif 'Heading 2' in style_name or (len(text) < 80 and p.runs and p.runs[0].bold and any(text.startswith(x) for x in ['Case ', '1.', '2.', '3.', '4.', '5.', 'Section', 'III.', 'IV.', 'V.'])):
                body_html.append(f'<h2>{html.escape(text)}</h2>')
            elif 'Heading 3' in style_name:
                body_html.append(f'<h3>{html.escape(text)}</h3>')
            elif style_name.startswith('List') or text.startswith('•') or text.startswith('- '):
                clean_text = text.lstrip('•- ').strip()
                body_html.append(f'<ul><li>{html.escape(clean_text)}</li></ul>')
            elif text.lower().startswith('figure') or text.lower().startswith('fig.'):
                body_html.append(f'<p class="fig-caption">{html.escape(text)}</p>')
            else:
                formatted = format_para(p)
                if formatted: body_html.append(f'<p>{formatted}</p>')
        elif isinstance(child, CT_Tbl):
            t = Table(child, doc)
            table_html = ['<div class="table-wrapper"><table>']
            for row_idx, row in enumerate(t.rows):
                table_html.append('<tr>')
                for cell in row.cells:
                    tag = 'th' if row_idx == 0 else 'td'
                    cell_text = html.escape(cell.text.strip()).replace('\n', '<br>')
                    table_html.append(f'<{tag}>{cell_text}</{tag}>')
                table_html.append('</tr>')
            table_html.append('</table></div>')
            body_html.append(''.join(table_html))

    doc_name_clean = filename.replace('.docx', '').replace('_', ' ')
    out_html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{html.escape(doc_name_clean)} - Document Viewer</title>
<style>{COMMON_CSS}</style>
</head>
<body>
<div class="preview-container">
  <div class="doc-header">
    <span class="badge-pill">Word Document (.docx)</span>
    <h1 class="doc-title">{html.escape(doc_name_clean)}</h1>
    <div class="doc-meta">
      <div class="meta-item">Author: <strong>Onkar Shelke</strong> (25102C2002)</div>
      <div class="meta-item">Semester: <strong>5 (TE CMPN) · VIT Mumbai</strong></div>
    </div>
  </div>
  <div class="actions-bar">
    <span style="font-size:12px;color:#64748b;">Direct in-dashboard preview · No download needed</span>
    <div>
      <a class="btn-action" href="../assignment-files/{html.escape(filename)}" download>Download Original</a>
      <button class="btn-action" onclick="window.print()">Print / PDF</button>
    </div>
  </div>
  <div class="doc-body">{''.join(body_html)}</div>
</div>
</body>
</html>"""
    out_file = os.path.join(PREVIEWS_DIR, f"{filename}.html")
    with open(out_file, 'w', encoding='utf-8') as f:
        f.write(out_html)
    print(f"DOCX preview: {filename}")

def convert_xlsx(filename):
    full_path = os.path.join(BASE_DIR, filename)
    if not os.path.exists(full_path): return
    wb = openpyxl.load_workbook(full_path, data_only=True)
    tabs_html = []
    sheets_html = []
    
    for i, sname in enumerate(wb.sheetnames):
        ws = wb[sname]
        active_cls = " active" if i == 0 else ""
        tabs_html.append(f'<button class="sheet-tab{active_cls}" onclick="switchSheet({i})">{html.escape(sname)}</button>')
        
        # Build table
        sheet_rows = []
        max_rows = min(ws.max_row, 150)
        max_cols = min(ws.max_column, 25)
        
        has_content = False
        table_html = ['<div class="table-wrapper"><table>']
        for r_idx in range(1, max_rows + 1):
            row_cells = [ws.cell(r_idx, c_idx).value for c_idx in range(1, max_cols + 1)]
            if not any(v is not None and str(v).strip() != '' for v in row_cells):
                continue
            has_content = True
            is_header = (r_idx == 1) or (r_idx == 2 and not any(row_cells[:3]))
            tag = 'th' if is_header else 'td'
            table_html.append('<tr>')
            for val in row_cells:
                v_str = '' if val is None else str(val)
                table_html.append(f'<{tag}>{html.escape(v_str)}</{tag}>')
            table_html.append('</tr>')
        table_html.append('</table></div>')
        
        if not has_content:
            table_html = ['<p style="color:#64748b;padding:20px;text-align:center;">Empty sheet or formula reference.</p>']
            
        sheets_html.append(f'<div id="sheet_{i}" class="sheet-content{active_cls}"><h3 style="margin-bottom:12px;color:#1e293b;">Sheet: {html.escape(sname)}</h3>{"".join(table_html)}</div>')
        
    doc_name_clean = filename.replace('.xlsx', '').replace('_', ' ')
    out_html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{html.escape(doc_name_clean)} - Spreadsheet Viewer</title>
<style>
{COMMON_CSS}
</style>
<script>
function switchSheet(idx) {{
  document.querySelectorAll('.sheet-tab').forEach((t, i) => t.classList.toggle('active', i === idx));
  document.querySelectorAll('.sheet-content').forEach((c, i) => c.classList.toggle('active', i === idx));
}}
</script>
</head>
<body>
<div class="preview-container">
  <div class="doc-header">
    <span class="badge-pill" style="background:#ecfdf5;color:#047857;">Excel Workbook (.xlsx)</span>
    <h1 class="doc-title">{html.escape(doc_name_clean)}</h1>
    <div class="doc-meta">
      <div class="meta-item">Author: <strong>Onkar Shelke</strong> (25102C2002)</div>
      <div class="meta-item">Sheets: <strong>{len(wb.sheetnames)}</strong></div>
      <div class="meta-item">Semester: <strong>5 · VIT Mumbai</strong></div>
    </div>
  </div>
  <div class="actions-bar">
    <span style="font-size:12px;color:#64748b;">Direct in-dashboard spreadsheet preview · No download needed</span>
    <a class="btn-action" href="../assignment-files/{html.escape(filename)}" download>Download Workbook</a>
  </div>
  <div class="sheet-tabs">
    {''.join(tabs_html)}
  </div>
  {''.join(sheets_html)}
</div>
</body>
</html>"""
    out_file = os.path.join(PREVIEWS_DIR, f"{filename}.html")
    with open(out_file, 'w', encoding='utf-8') as f:
        f.write(out_html)
    print(f"XLSX preview: {filename}")

def convert_pptx(filename):
    full_path = os.path.join(BASE_DIR, filename)
    if not os.path.exists(full_path): return
    with zipfile.ZipFile(full_path) as z:
        slides = sorted([f for f in z.namelist() if f.startswith('ppt/slides/slide') and f.endswith('.xml')])
        slides_data = []
        for s in slides:
            xml_content = z.read(s)
            tree = ET.fromstring(xml_content)
            texts = [node.text.strip() for node in tree.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}t') if node.text and node.text.strip()]
            slides_data.append(texts)
            
    # Generate presentation slides HTML
    slides_rendered = []
    for i, texts in enumerate(slides_data):
        active = " active" if i == 0 else ""
        title = texts[0] if len(texts) > 0 else f"Slide {i+1}"
        subtitle = texts[1] if len(texts) > 1 else ""
        body_bullets = texts[2:] if len(texts) > 2 else []
        
        bullets_html = "".join([f'<li style="margin-bottom:8px;font-size:15px;color:#334155;">{html.escape(b)}</li>' for b in body_bullets])
        
        slides_rendered.append(f"""
        <div class="slide-card{active}" id="slide_{i}" style="min-height:360px;background:#ffffff;border:1px solid #e2e8f0;border-radius:12px;padding:32px;display:none;flex-direction:column;justify-content:center;box-shadow:0 4px 14px rgba(0,0,0,0.03);">
          <div style="display:flex;justify-content:space-between;align-items:center;margin-bottom:14px;">
            <span style="font-size:11px;font-weight:700;color:#2563eb;text-transform:uppercase;letter-spacing:1px;">Slide {i+1} of {len(slides_data)}</span>
            <span style="font-size:12px;color:#64748b;">E-Waste to Smart Products Pitch Deck</span>
          </div>
          <h2 style="font-size:24px;font-weight:800;color:#0f172a;margin-bottom:8px;">{html.escape(title)}</h2>
          {f'<p style="font-size:16px;color:#475569;margin-bottom:16px;font-weight:500;">{html.escape(subtitle)}</p>' if subtitle else ''}
          {f'<ul style="margin-left:20px;margin-top:10px;">{bullets_html}</ul>' if bullets_html else ''}
        </div>
        """)
        
    doc_name_clean = filename.replace('.pptx', '').replace('_', ' ')
    out_html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{html.escape(doc_name_clean)} - Presentation Viewer</title>
<style>
{COMMON_CSS}
.slide-card.active {{ display: flex !important; }}
.nav-btn {{
  background: var(--primary);
  color: white;
  border: none;
  padding: 8px 18px;
  border-radius: 8px;
  font-weight: 700;
  cursor: pointer;
}}
.nav-btn:disabled {{
  background: #cbd5e1;
  cursor: not-allowed;
}}
</style>
<script>
let curSlide = 0;
const total = {len(slides_data)};
function showSlide(idx) {{
  if (idx < 0 || idx >= total) return;
  curSlide = idx;
  document.querySelectorAll('.slide-card').forEach((s, i) => s.classList.toggle('active', i === curSlide));
  document.getElementById('prevBtn').disabled = curSlide === 0;
  document.getElementById('nextBtn').disabled = curSlide === total - 1;
  document.getElementById('slideIndicator').textContent = (curSlide + 1) + ' / ' + total;
}}
document.addEventListener('keydown', e => {{
  if (e.key === 'ArrowRight' || e.key === ' ') showSlide(curSlide + 1);
  if (e.key === 'ArrowLeft') showSlide(curSlide - 1);
}});
</script>
</head>
<body>
<div class="preview-container">
  <div class="doc-header">
    <span class="badge-pill" style="background:#fff7ed;color:#c2410c;">Pitch Deck (.pptx)</span>
    <h1 class="doc-title">{html.escape(doc_name_clean)}</h1>
    <div class="doc-meta">
      <div class="meta-item">Author: <strong>Onkar Shelke</strong> (25102C2002)</div>
      <div class="meta-item">Topic: <strong>Novel Idea Pitch Deck</strong></div>
    </div>
  </div>
  <div class="actions-bar">
    <span style="font-size:12px;color:#64748b;">Direct slide presentation · Use Left / Right keys</span>
    <a class="btn-action" href="../assignment-files/{html.escape(filename)}" download>Download PPTX</a>
  </div>
  <div class="doc-body">
    {''.join(slides_rendered)}
    <div style="display:flex;justify-content:space-between;align-items:center;margin-top:20px;">
      <button id="prevBtn" class="nav-btn" disabled onclick="showSlide(curSlide - 1)">← Previous</button>
      <span id="slideIndicator" style="font-weight:700;color:#1e293b;">1 / {len(slides_data)}</span>
      <button id="nextBtn" class="nav-btn" onclick="showSlide(curSlide + 1)">Next →</button>
    </div>
  </div>
</div>
</body>
</html>"""
    out_file = os.path.join(PREVIEWS_DIR, f"{filename}.html")
    with open(out_file, 'w', encoding='utf-8') as f:
        f.write(out_html)
    print(f"PPTX preview: {filename}")

# Run for all
for f in sorted(os.listdir(BASE_DIR)):
    ext = os.path.splitext(f)[1].lower()
    if ext == '.docx': convert_docx(f)
    elif ext == '.xlsx': convert_xlsx(f)
    elif ext == '.pptx': convert_pptx(f)

print("All file conversions complete!")
